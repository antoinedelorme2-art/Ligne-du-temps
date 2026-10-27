# -*- coding: utf-8 -*-
"""Régénère donnees.json à partir du classeur.

    python generer_donnees.py Base_lignes_du_temps.xlsx
    python generer_donnees.py <id-du-google-sheets> --sortie site/donnees.json
    python generer_donnees.py https://docs.google.com/spreadsheets/d/<id>/edit

Dépendance unique : pip install openpyxl

Le classeur reste la source de vérité. Ce script ne fait que traduire ses
onglets « Lignes du temps » et « Objets » dans le format compact que la page
sait lire. Un classeur Google est téléchargé au format xlsx ; il doit être
partagé en lecture par lien, sans quoi Google renvoie une page de connexion
au lieu du fichier.

S'y ajoutent les banques ouvertes rangées dans sources/ par
recolter_sources.py (Wikidata, chronologies Wikipédia, PeriodO). Les objets
qui décrivent le même événement d'une banque à l'autre sont rapprochés (même
identifiant Wikidata, ou même intitulé à la même époque dans la même région),
et chaque événement vu par plusieurs banques reçoit un objet dans la banque
« Synthèse ».
"""
import argparse
import json
import os
import re
import sys
import tempfile
import urllib.error
import urllib.request
from datetime import date

from openpyxl import load_workbook

RE_SHEETS = re.compile(r"docs\.google\.com/spreadsheets/d/([\w-]{20,})")
RE_DRIVE = re.compile(r"drive\.google\.com/file/d/([\w-]{20,})")
RE_ID = re.compile(r"^[\w-]{30,}$")


def echec(*lignes):
    """Message d'erreur visible dans l'interface d'Actions, puis arrêt."""
    print("::error::" + lignes[0])
    for l in lignes[1:]:
        print("   " + l)
    sys.exit(1)


def essayer(url):
    """Télécharge une URL. Renvoie (octets, statut, message) sans lever."""
    req = urllib.request.Request(url, headers={"User-Agent": "Mozilla/5.0"})
    try:
        with urllib.request.urlopen(req, timeout=180) as r:
            return r.read(), r.status, r.headers.get("Content-Type", "")
    except urllib.error.HTTPError as e:
        return b"", e.code, str(e)
    except Exception as e:                                  # réseau, DNS, délai
        return b"", 0, str(e)


def obtenir(source):
    """Renvoie un chemin local, en téléchargeant d'abord si nécessaire."""
    source = (source or "").strip()
    ident = None
    m = RE_SHEETS.search(source)
    if m:
        ident = m.group(1)
    elif RE_DRIVE.search(source):
        ident = RE_DRIVE.search(source).group(1)
    elif RE_ID.match(source):
        ident = source
    if not ident:
        if not os.path.exists(source):
            echec(f"Classeur introuvable : {source!r}",
                  "Attendu : un chemin de fichier, un identifiant Google, ou une URL.")
        return source

    # Deux points d'entrée : un vrai Google Sheets, ou un .xlsx déposé dans
    # Drive sans conversion. Le second ne répond pas à l'export Sheets.
    tentatives = [
        ("classeur Google Sheets",
         f"https://docs.google.com/spreadsheets/d/{ident}/export?format=xlsx"),
        ("fichier .xlsx dans Drive",
         f"https://drive.google.com/uc?export=download&id={ident}"),
    ]
    dernier = ""
    for quoi, url in tentatives:
        print(f"Tentative : {quoi}…")
        donnees, statut, info = essayer(url)
        if donnees[:2] == b"PK":
            tmp = tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False)
            tmp.write(donnees)
            tmp.close()
            print(f"  reçu : {len(donnees) / 1e6:.1f} Mo")
            return tmp.name
        apercu = donnees[:200].decode("utf-8", "replace").replace("\n", " ")
        dernier = f"HTTP {statut}, {info}, début de la réponse : {apercu!r}"
        print(f"  échec — {dernier}")

    echec(
        "Impossible de récupérer le classeur.",
        f"Identifiant utilisé : {ident}",
        f"Dernière réponse : {dernier}",
        "",
        "Trois causes possibles, par ordre de fréquence :",
        "1. Le fichier n'est pas partagé. Ouvrez-le, Partager → Accès général →",
        "   « Tous les utilisateurs disposant du lien », rôle Lecteur.",
        "2. C'est un .xlsx déposé dans Drive, pas un Google Sheets. Ouvrez-le et",
        "   faites Fichier → Enregistrer au format Google Sheets, puis reprenez",
        "   l'identifiant du NOUVEAU document.",
        "3. L'identifiant du secret CLASSEUR est celui d'un dossier Drive et non",
        "   d'un document.",
    )


ap = argparse.ArgumentParser()
ap.add_argument("source", nargs="?", default="Base_lignes_du_temps.xlsx",
                help="chemin du .xlsx, ou identifiant / URL du Google Sheets")
ap.add_argument("--sortie", default="donnees.json")
args = ap.parse_args()

SRC = obtenir(args.source)
DST = args.sortie

PETITS = {"and", "of", "the", "in", "for", "&"}
SIGLES = {"UK", "USA"}


def titre_en(s):
    """« WARS AND BATTLES » → « Wars and Battles », sigles préservés."""
    mots = str(s).replace("\n", " ").split()
    out = []
    for i, w in enumerate(mots):
        if w in SIGLES or not w.isalpha():
            out.append(w if w in SIGLES or not w.isupper() else w.capitalize())
        elif i > 0 and w.lower() in PETITS:
            out.append(w.lower())
        else:
            out.append(w.capitalize() if w.isupper() else w)
    return " ".join(out)


def minutes_dans_annee(an, mois, jour, heure):
    """Position dans l'année, en minutes depuis le 1er janvier 00 h 00.

    0 signifie « inconnue » ; le 1er janvier vaut donc 1440. Seules les dates
    de l'ère chrétienne sont converties, la source n'ayant jamais de mois
    avant l'an 1.
    """
    if an is None or an < 1 or not mois:
        return 0
    try:
        j = date(int(an), int(mois), int(jour) if jour else 1)
    except ValueError:
        return 0
    mn = 0
    if heure not in (None, ""):
        m = re.match(r"^(\d{1,2})[:h.]?(\d{2})?", str(heure).strip())
        if m:
            mn = int(m.group(1)) * 60 + int(m.group(2) or 0)
    return j.timetuple().tm_yday * 1440 + mn


wb = load_workbook(SRC, read_only=True, data_only=True)

for onglet in ("Lignes du temps", "Objets"):
    if onglet not in wb.sheetnames:
        echec(f"Onglet « {onglet} » absent du classeur.",
              "Onglets trouvés : " + ", ".join(wb.sheetnames))


def verifier(ws, attendues):
    presentes = [c.value for c in next(ws.iter_rows(max_row=1))]
    manquantes = [c for c in attendues if c not in presentes]
    if manquantes:
        echec(f"Colonnes manquantes dans l'onglet « {ws.title} » : "
              + ", ".join(manquantes),
              "En-têtes trouvés : " + ", ".join(str(p) for p in presentes if p))
    return {h: i for i, h in enumerate(presentes)}


# ------------------------------------------------------------ lignes du temps
ws = wb["Lignes du temps"]
verifier(ws, ["ID", "Nom (clé)", "Nom français", "Section", "Section (fr)", "Couleur"])
idx = verifier(ws, ["ID", "Nom (clé)", "Section"])
lignes, cle_vers_i = [], {}
sections, sec_index, sec_fr_index = [], {}, {}

for row in ws.iter_rows(min_row=2, values_only=True):
    if not row or not row[idx["Nom (clé)"]]:
        continue
    cle = row[idx["Nom (clé)"]]
    sec = row[idx["Section"]]
    sec_fr = row[idx["Section (fr)"]] or sec
    coul = row[idx["Couleur"]] or "#7F8C8D"
    if sec not in sec_index:
        sec_index[sec] = len(sections)
        sec_fr_index[sec_fr] = len(sections)
        sections.append([sec_fr, titre_en(sec), coul])
    cle_vers_i[cle] = len(lignes)
    lignes.append([row[idx["ID"]], cle, row[idx["Nom français"]] or "",
                   sec, sec_fr, coul])


def ensembles(txt):
    """« Égypte ; Syrie » → indices de sections ; 0 quand il n'y a rien."""
    if not txt:
        return 0
    i = [sec_fr_index[p.strip()] for p in str(txt).split(";")
         if p.strip() in sec_fr_index]
    return i if len(i) > 1 else 0


# ------------------------------------------------------------------ banques
# La banque « Climat et catastrophes » ne donne pas d'objets : ses lignes
# deviennent les repères du bandeau fixe en haut de la page.
BANQUE_CLIMAT = "Climat et catastrophes"
banques = []
if "Listes" in wb.sheetnames:
    ws = wb["Listes"]
    entetes = [c.value for c in next(ws.iter_rows(max_row=1))]
    if "Banque" in entetes:
        col = entetes.index("Banque")
        for row in ws.iter_rows(min_row=2, values_only=True):
            b = row[col] if col < len(row) else None
            if b and b != BANQUE_CLIMAT and b not in banques:
                banques.append(b)


def indice_banque(b):
    b = b or "—"
    if b not in banques:
        banques.append(b)
    return banques.index(b)


# Noms anglais et couleurs de repli des catégories du bandeau climatique.
CATS_EN = {
    "Épisodes climatiques": ("Climatic episodes", "#2C6E8B"),
    "Séismes et tsunamis": ("Earthquakes and tsunamis", "#6E4B1F"),
    "Sécheresses et famines": ("Droughts and famines", "#B5871D"),
    "Éruptions volcaniques": ("Volcanic eruptions", "#8E2317"),
    "Épidémies": ("Epidemics", "#5B3A6B"),
    "Impacts et événements cosmiques": ("Cosmic events", "#3F4A52"),
}
categories, cat_index = [], {}


def indice_categorie(tl):
    """« Climat – Épidémies » → indice de la catégorie « Épidémies »."""
    nom = re.sub(r"^Climat\s*[–-]\s*", "", str(tl or "")).strip() or "Autre"
    if nom not in cat_index:
        en, coul = CATS_EN.get(nom, (nom, "#3F4A52"))
        if tl in cle_vers_i:
            coul = lignes[cle_vers_i[tl]][5] or coul
        cat_index[nom] = len(categories)
        categories.append([nom, en, coul])
    return cat_index[nom]


# Rapprochements entre banques (même événement vu par deux sources). Lus
# dans une colonne « Groupe » si le classeur en a une ; sinon repris du
# donnees.json précédent, où ils ont été établis.
GROUPE_AUTO = 100000             # au-delà : groupes recalculés à chaque fois
groupes_precedents = {}
climat_precedent = None          # (categories, reperes) du donnees.json versé
for chemin in ("donnees.json", DST):
    try:
        with open(chemin, encoding="utf-8") as f:
            ancien = json.load(f)
        if climat_precedent is None and ancien.get("reperes"):
            climat_precedent = (ancien.get("categories", []), ancien["reperes"])
        for o in ancien.get("objets", []):
            if len(o) > 17 and 0 < o[17] < GROUPE_AUTO:
                cle_l = ancien["lignes"][o[2]][1]
                groupes_precedents[(o[0], cle_l, ancien["banques"][o[16]], o[3])] = o[17]
        if groupes_precedents:
            break
    except (OSError, ValueError, KeyError, IndexError, TypeError):
        continue


# -------------------------------------------------------------------- objets
ws = wb["Objets"]
k = verifier(ws, ["Objet", "Version française", "Timeline", "Début (num)",
                  "Fin (num)", "Précision début", "Qualificatif début",
                  "Marge début (± ans)", "Marge fin (± ans)", "Notes", "URL1",
                  "Fiabilité date", "Sous-piste", "Ensembles liés",
                  "Start Month", "Start Date", "Start Hour",
                  "End Month", "End Date", "End Hour", "Banque"])
k_ampleur = k.get("Ampleur")
k_liens = [k[c] for c in ("URL2", "URL3") if c in k]
k_groupe = k.get("Groupe")
objets, reperes, orphelins = [], [], set()

for row in ws.iter_rows(min_row=2, values_only=True):
    if not row or not row[k["Objet"]]:
        continue
    tl = row[k["Timeline"]]
    banque = row[k["Banque"]]
    debut = row[k["Début (num)"]]
    if debut in (None, ""):
        continue
    fin = row[k["Fin (num)"]]
    fin = int(fin) if isinstance(fin, (int, float)) else None
    if fin is not None and fin < int(debut):
        fin = None
    notes = (row[k["Notes"]] or "").strip()
    if len(notes) > 900:
        notes = notes[:900].rsplit(" ", 1)[0] + "…"
    if banque == BANQUE_CLIMAT:
        ampleur = row[k_ampleur] if k_ampleur is not None else None
        reperes.append([
            row[k["Objet"]],
            row[k["Version française"]] or "",
            indice_categorie(tl),
            int(ampleur) if isinstance(ampleur, (int, float)) else 1,
            int(debut),
            fin,
            notes,
            row[k["URL1"]] or "",
        ])
        continue
    if tl not in cle_vers_i:
        orphelins.add(tl)
        continue
    groupe = row[k_groupe] if k_groupe is not None else None
    if not isinstance(groupe, (int, float)) or not groupe:
        groupe = groupes_precedents.get(
            (row[k["Objet"]], tl, banque or "—", int(debut)), 0)
    objets.append([
        row[k["Objet"]],
        row[k["Version française"]] or "",
        cle_vers_i[tl],
        int(debut),
        fin,
        row[k["Précision début"]] or "année",
        row[k["Qualificatif début"]] or "",
        row[k["Marge début (± ans)"]] or 0,
        row[k["Marge fin (± ans)"]] or 0,
        notes,
        row[k["URL1"]] or "",
        1 if row[k["Fiabilité date"]] in ("texte", "vérifiée", "corrigée") else 0,
        row[k["Sous-piste"]] or "",
        ensembles(row[k["Ensembles liés"]]),
        minutes_dans_annee(debut, row[k["Start Month"]], row[k["Start Date"]],
                           row[k["Start Hour"]]),
        minutes_dans_annee(fin, row[k["End Month"]], row[k["End Date"]],
                           row[k["End Hour"]]),
        indice_banque(banque),
        int(groupe),
        [row[c] for c in k_liens if c < len(row) and row[c]] or 0,
        0,
    ])

# ---------------------------------------------------------- banques ouvertes
# Les fichiers de sources/ sont produits par recolter_sources.py (workflow
# « Récolter les banques ouvertes »). Chaque objet y porte déjà sa piste.
SOURCES = ("wikidata", "wikidata_regnes", "wikipedia", "periodo")
BANQUE_SYNTHESE = "Synthèse"
RANG_PRECISION = {"jour": 0, "mois": 1, "année": 2, "décennie": 3,
                  "siècle": 4, "millénaire": 5}
# Pistes thématiques : un objet qui s'y trouve peut correspondre à un objet
# rangé sous son pays dans une autre banque.
PISTES_GENERIQUES = {"Wars And Battles",
                     "Other Events – Other Events and Developments",
                     "European Monarchies",
                     "European Monarchies – European Monarchies and Renaissance Rulers",
                     "Rome And Persia – Rulers"}
TYPES_DURABLES = {"État historique", "Pandémie", "Épidémie"}
MOTS_AUTRE_EVENEMENT = {"died", "dies", "dead", "death", "born", "killed",
                        "assassinated", "executed", "deposed", "meurt", "mort",
                        "nait", "naissance", "assassine", "execute", "depose"}
vies = {}                        # indice d'objet → « Naissance : … · Mort : … »
mentions = {}                    # indice d'objet → QID des articles cités


def section_de_cle(cle):
    return sec_fr_index[lignes[cle_vers_i[cle]][4]]


licences = {}


# Pas d'objet des banques ouvertes avant le plus ancien repère climatique du
# bandeau : il fixe le début de la frise. Au-delà, ce sont des mythes (règne
# d'Amaterasu, « 2 343 127 av. J.-C. ») ou de la préhistoire qui étirerait
# l'axe et la barre de défilement. Une période de plus de 4 000 ans
# (« Holocene », « Prehistoric, unspecified ») n'apprend rien non plus.
DUREE_MAX = 4000
ANNEE = date.today().year


def charger_sources():
    ici = os.path.dirname(os.path.abspath(__file__))
    plancher = min((r[4] for r in reperes), default=-6500)
    ecartes = 0
    qids = {}                    # indice d'objet → identifiant Wikidata
    for nom in SOURCES:
        chemin = os.path.join(ici, "sources", nom + ".json")
        try:
            with open(chemin, encoding="utf-8") as f:
                src = json.load(f)
        except (OSError, ValueError) as e:
            print(f"Banque ouverte « {nom} » absente ou illisible ({e}) : ignorée.")
            continue
        b = indice_banque(src["banque"])
        licences[src["banque"]] = src.get("licence", "")
        sans_piste = 0
        for s in src["objets"]:
            if s.get("piste") not in cle_vers_i:
                sans_piste += 1
                continue
            secs = sorted({section_de_cle(c) for c in s.get("ensembles", [])
                           if c in cle_vers_i})
            debut, fin = int(s["debut"]), s.get("fin")
            fin = int(fin) if isinstance(fin, (int, float)) and fin >= debut else None
            # Date connue seulement au siècle : le début de ce siècle (Wikidata
            # note « XXIe siècle » 2100, ce qui placerait l'objet dans le futur).
            if s.get("precision") in ("siècle", "millénaire"):
                n = (abs(debut) - 1) // 100
                debut = n * 100 + 1 if debut > 0 else -(n + 1) * 100
            # Fin dans le futur (« jusqu'en 2100 ») : arrêtée à aujourd'hui.
            if fin is not None and fin > ANNEE:
                fin = ANNEE
            if (debut < plancher or debut > ANNEE
                    or (fin is not None and fin - debut > DUREE_MAX)):
                ecartes += 1
                continue
            notes = (s.get("notes") or "").strip()
            if s.get("qid"):
                qids[len(objets)] = s["qid"]
            if s.get("vie"):
                vies[len(objets)] = s["vie"]
            if s.get("mentions"):
                mentions[len(objets)] = s["mentions"]
            # nature : « r » règne, « o » office ; « j » date du calendrier julien
            nature = (s.get("nature") or "") + ("j" if s.get("julien") else "")
            objets.append([
                s["nom"], s.get("fr") or "", cle_vers_i[s["piste"]], debut, fin,
                s.get("precision") or "année", s.get("qual") or "", 0, 0,
                notes[:900], s.get("url") or "", 1, s.get("type") or "",
                secs if len(secs) > 1 else 0,
                minutes_dans_annee(debut, s.get("mois"), s.get("jour"), None),
                minutes_dans_annee(fin, s.get("moisF"), s.get("jourF"), None),
                b, 0, s.get("liens") or 0, nature or 0,
            ])
        print(f"Banque ouverte {src['banque']} : {len(src['objets']) - sans_piste} "
              f"objets (récolte du {src.get('recolte', '?')})"
              + (f", {sans_piste} sans piste connue" if sans_piste else ""))
    if ecartes:
        print(f"Banques ouvertes : {ecartes} objets écartés (avant {-plancher} av. J.-C., "
              f"après {ANNEE} ou période de plus de {DUREE_MAX} ans)")
    return qids


def cle_titre(t):
    from rattachement import cle_titre as ct
    return ct(t)


def tolerance(o):
    g = {"siècle": 50, "décennie": 5, "millénaire": 300}.get(o[5], 1)
    if o[6] == "circa":
        g = max(g, 10)
    return max(g, int(o[7] or 0))


RE_DATES_FIN = re.compile(
    r"(?:\s*\([^()]*\d{3}[^()]*\)|,\s*(?:c\.\s*)?-?\d{1,4}\s*(?:BC|BCE|AD|CE)?"
    r"(?:\s*[-–]\s*-?\d{1,4}\s*(?:BC|BCE|AD|CE)?)?)\s*$")


def sans_dates(t):
    """« Louis VII, 1137-1180 » ou « Old Kingdom (2686–2181 BC) » → sans dates."""
    court = RE_DATES_FIN.sub("", t or "").strip()
    return court if len(court) >= 3 else t


def fusionner_sources():
    """Ajoute les banques ouvertes, rapproche les doublons, crée « Synthèse »."""
    qids = charger_sources()
    n = len(objets)
    parent = list(range(n))
    taille = [1] * n

    def racine(i):
        while parent[i] != i:
            parent[i] = parent[parent[i]]
            i = parent[i]
        return i

    def unir(a, b, plafond=12):
        ra, rb = racine(a), racine(b)
        if ra == rb or taille[ra] + taille[rb] > plafond:
            return False
        if taille[ra] < taille[rb]:
            ra, rb = rb, ra
        parent[rb] = ra
        taille[ra] += taille[rb]
        return True

    # rapprochements établis à la main : conservés tels quels
    manuels = {}
    for i, o in enumerate(objets):
        if o[17]:
            if o[17] in manuels:
                unir(manuels[o[17]], i, plafond=10 ** 6)
            else:
                manuels[o[17]] = i
    # même identifiant Wikidata
    par_qid = {}
    for i, q in qids.items():
        if q in par_qid:
            unir(par_qid[q], i)
        else:
            par_qid[q] = i
    def compatibles(a, b):
        """Même époque, et même région ou piste thématique."""
        oa, ob = objets[a], objets[b]
        if abs(oa[3] - ob[3]) > max(tolerance(oa), tolerance(ob)):
            return False
        if oa[16] == ob[16] and oa[2] != ob[2]:
            return False                  # même banque, autre piste : distinct
        sa = {lignes[oa[2]][4]} | ({sections[k][0] for k in oa[13]} if oa[13] else set())
        sb = {lignes[ob[2]][4]} | ({sections[k][0] for k in ob[13]} if ob[13] else set())
        return bool(sa & sb or lignes[oa[2]][1] in PISTES_GENERIQUES
                    or lignes[ob[2]][1] in PISTES_GENERIQUES)

    def mots_forts(i):
        """Mots porteurs de l'intitulé (4 lettres et plus, ou nombres)."""
        return {m for t in (objets[i][0], objets[i][1]) for m in cle_titre(t).split()
                if len(m) >= 4 or m.isdigit()}

    def personne(i):
        return bool(objets[i][19] and objets[i][19].strip("j"))

    # une ligne de chronologie qui cite l'article d'un événement de la même année
    # (les deux premiers articles seulement : le sujet de la ligne ; jamais un
    # État ou une épidémie, souvent cités en passant)
    for i, qs in mentions.items():
        for q in qs[:2]:
            j = par_qid.get(q)
            if (j is not None and objets[j][12] not in TYPES_DURABLES
                    and abs(objets[i][3] - objets[j][3]) <= 1
                    and len(mots_forts(i) & mots_forts(j)) >= 2):
                unir(i, j)
    # même intitulé (anglais ou français), même époque, même région
    cles = {}
    par_cle = {}
    for i, o in enumerate(objets):
        cles[i] = {cle_titre(o[0]), cle_titre(o[1])} - {""}
        for t in cles[i]:
            if len(t) >= 8 or (" " in t and len(t) >= 5):
                par_cle.setdefault(t, []).append(i)
    for t, membres in par_cle.items():
        if len(membres) < 2 or len(membres) > 40:
            continue                      # intitulé trop générique
        for x in range(len(membres)):
            for y in range(x + 1, len(membres)):
                if compatibles(membres[x], membres[y]):
                    unir(membres[x], membres[y])
    # une phrase qui commence par l'intitulé d'un événement :
    # « Battle of Stoke, the decisive engagement… » ↔ « Battle of Stoke »
    courts = {t: m for t, m in par_cle.items()
              if 2 <= t.count(" ") + 1 <= 6 and len(m) <= 4}
    for i in range(n):
        for t in cles[i]:
            mots = t.split()
            for k in range(min(6, len(mots) - 1), 1, -1):
                if mots[k] in MOTS_AUTRE_EVENEMENT:
                    continue              # « Chlothar III died » n'est pas son règne
                for j in courts.get(" ".join(mots[:k]), ()):
                    if j != i and not personne(j) and compatibles(i, j):
                        unir(i, j)

    groupes = {}
    for i in range(n):
        groupes.setdefault(racine(i), []).append(i)
    b_synth = indice_banque(BANQUE_SYNTHESE)
    # banques du classeur d'abord : ce sont les choix de l'auteur
    ordre_banque = {b: k for k, b in enumerate(banques)}
    prochain = GROUPE_AUTO
    nb_synth = 0
    for membres in groupes.values():
        if len(membres) < 2:
            continue
        numero = min((objets[i][17] for i in membres if objets[i][17]), default=0)
        if not numero:
            numero = prochain
            prochain += 1
        for i in membres:
            objets[i][17] = numero
        if len({objets[i][16] for i in membres}) < 2:
            continue                      # doublon interne à une banque
        # un règne ou un mandat l'emporte sur un simple point daté
        meilleur = min(membres, key=lambda i: (
            not (objets[i][19] and objets[i][19].strip("j")),
            -objets[i][11], RANG_PRECISION.get(objets[i][5], 2),
            ordre_banque[banques[objets[i][16]]]))
        m = objets[meilleur]
        fr = next((objets[i][1] for i in sorted(membres, key=lambda i: (
            banques[objets[i][16]] != "Wikidata", ordre_banque[banques[objets[i][16]]]))
            if objets[i][1]), "")
        secs = {lignes[objets[i][2]][4] for i in membres}
        for i in membres:
            if objets[i][13]:
                secs |= {sections[k][0] for k in objets[i][13]}
        ens = sorted(sec_fr_index[s] for s in secs)[:6]
        objets.append([sans_dates(m[0]), sans_dates(fr), m[2], m[3], m[4], m[5], m[6], m[7], m[8], "", "",
                       max(objets[i][11] for i in membres), m[12],
                       ens if len(ens) > 1 else 0, m[14], m[15], b_synth,
                       numero, 0, m[19]])
        nb_synth += 1
        # naissance et mort de la personne : dans la notice de chaque source
        vie = next((vies[i] for i in membres if i in vies), "")
        if vie:
            for i in membres:
                if i not in vies and vie not in objets[i][9]:
                    objets[i][9] = (vie + "\n\n" + objets[i][9]).strip()
    print(f"Synthèse : {nb_synth} événements vus par plusieurs banques "
          f"({sum(len(g) for g in groupes.values() if len(g) > 1)} objets regroupés)")


# Un classeur sans lignes « Climat et catastrophes » viderait le bandeau :
# on garde alors les repères du donnees.json versé dans le dépôt.
if not reperes and climat_precedent:
    categories, reperes = climat_precedent
    print(f"Aucun repère climatique dans le classeur : {len(reperes)} repris "
          "du donnees.json du dépôt.")

if not objets:
    echec("Aucun objet exploitable dans le classeur.",
          "Vérifiez que la colonne « Début (num) » contient bien des valeurs "
          "calculées et non des formules vides.")

fusionner_sources()

dossier = os.path.dirname(DST)
if dossier:
    os.makedirs(dossier, exist_ok=True)
with open(DST, "w", encoding="utf-8") as f:
    json.dump({"lignes": lignes, "sections": sections, "banques": banques,
               "objets": objets, "categories": categories,
               "reperes": reperes, "licences": licences},
              f, ensure_ascii=False, separators=(",", ":"))

print(f"{DST} : {len(objets)} objets, {len(lignes)} pistes, "
      f"{len(sections)} sections, {len(reperes)} repères climatiques")
print("Banques : " + ", ".join(
    f"{b} ({sum(1 for o in objets if o[16] == i)})"
    for i, b in enumerate(banques)))
if orphelins:
    print("Pistes citées dans Objets mais absentes du registre :",
          ", ".join(sorted(orphelins)[:10]))
